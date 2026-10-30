# Autor: Thiago Paolo Icochea Rodriguez (U22330428) · Servicios Cloud 45104 · UTP
"""Genera la hoja de cálculo de la evaluación con fórmulas (cargas, makespan, cota inferior).
Uso: python3 -m scheduler.excel → resultados/scheduling_caso.xlsx"""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from .caso import JOBS

OUT = os.environ.get("OUT", "resultados")
ROJO = PatternFill("solid", fgColor="B50D30")
SUAVE = PatternFill("solid", fgColor="F6E7EA")
B = Border(*(Side(style="thin", color="BBBBBB"),) * 4)


def cab(ws, fila, textos):
    for c, t in enumerate(textos, 1):
        x = ws.cell(fila, c, t); x.font = Font(bold=True, color="FFFFFF"); x.fill = ROJO; x.border = B


def main():
    os.makedirs(OUT, exist_ok=True)
    wb = Workbook(); ws = wb.active; ws.title = "Datos"
    ws["A1"] = "Scheduling en la nube · Thiago Paolo Icochea Rodriguez (U22330428) · Servicios Cloud 45104"; ws["A1"].font = Font(bold=True, size=12)
    cab(ws, 3, ["Trabajo", "Duración (min)"])
    for i, (j, d) in enumerate(JOBS.items(), 4):
        ws.cell(i, 1, j).border = B; ws.cell(i, 2, d).border = B
    ws["A12"] = "Tiempo total"; ws["B12"] = "=SUM(B4:B11)"
    ws["A13"] = "N.° de VM"; ws["B13"] = 3
    ws["A14"] = "Carga media teórica"; ws["B14"] = "=B12/B13"; ws["B14"].number_format = "0.00"
    ws["A15"] = "Cota inferior entera"; ws["B15"] = "=ROUNDUP(B14,0)"
    for r in range(12, 16): ws.cell(r, 1).font = Font(bold=True); ws.cell(r, 2).fill = SUAVE

    asign = {
        "Round Robin": {"VM1": ["J1", "J4", "J7"], "VM2": ["J2", "J5", "J8"], "VM3": ["J3", "J6"]},
        "SJF (VM libre primero)": {"VM1": ["J7", "J6", "J5"], "VM2": ["J4", "J8", "J1"], "VM3": ["J2", "J3"]},
        "Propuesta equilibrada": {"VM1": ["J1", "J6"], "VM2": ["J5", "J4", "J2"], "VM3": ["J3", "J8", "J7"]},
    }
    fila_job = {j: i for i, j in enumerate(JOBS, 4)}
    for nombre, plan in asign.items():
        h = wb.create_sheet(nombre[:28])
        h["A1"] = nombre; h["A1"].font = Font(bold=True, size=12)
        cab(h, 3, ["VM", "Trabajos asignados", "Carga total (min)"])
        for r, (vm, js) in enumerate(plan.items(), 4):
            h.cell(r, 1, vm); h.cell(r, 2, " + ".join(js))
            h.cell(r, 3, "=" + "+".join(f"Datos!B{fila_job[j]}" for j in js))
            for c in range(1, 4): h.cell(r, c).border = B
        h["A8"] = "Makespan Cmax"; h["C8"] = "=MAX(C4:C6)"
        h["A9"] = "Desbalance (máx − mín)"; h["C9"] = "=MAX(C4:C6)-MIN(C4:C6)"
        h["A10"] = "Utilización media (%)"; h["C10"] = "=SUM(C4:C6)/(3*C8)*100"; h["C10"].number_format = "0.0"
        h["A11"] = "Diferencia con la cota inferior"; h["C11"] = "=C8-Datos!B15"
        for r in range(8, 12): h.cell(r, 1).font = Font(bold=True); h.cell(r, 3).fill = SUAVE
        h.column_dimensions["A"].width = 30; h.column_dimensions["B"].width = 24; h.column_dimensions["C"].width = 18
    sjf = wb.create_sheet("Orden SJF")
    cab(sjf, 1, ["Posición", "Trabajo", "Duración"])
    for p, j in enumerate(sorted(JOBS, key=lambda x: (JOBS[x], x)), 1):
        sjf.cell(p + 1, 1, p); sjf.cell(p + 1, 2, j); sjf.cell(p + 1, 3, f"=Datos!B{fila_job[j]}")
    ws.column_dimensions["A"].width = 24; ws.column_dimensions["B"].width = 16
    wb.save(f"{OUT}/scheduling_caso.xlsx"); print(f"{OUT}/scheduling_caso.xlsx")


if __name__ == "__main__":
    main()
